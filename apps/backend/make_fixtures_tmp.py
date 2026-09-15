"""GEÇİCİ — Alan 9/10/11 test dosyalarını üretir. İş bitince silinecek."""
from __future__ import annotations

import zipfile
from pathlib import Path

OUT = Path("C:/Users/kerew/AppData/Local/Temp/stuhub_test_fixtures")
OUT.mkdir(parents=True, exist_ok=True)

TEXT = (
    "[TEST] Gelisim Psikolojisi Notu. Piaget'nin bilissel gelisim kurami dort "
    "evreden olusur: duyusal-motor, islem oncesi, somut islemler ve soyut islemler. "
    "Vygotsky ise yakinsak gelisim alanini (ZPD) vurgular. Baglanma kurami Bowlby "
    "tarafindan gelistirilmis, Ainsworth'un yabanci ortam deneyi ile sinanmistir."
)

made: list[str] = []

# ── PDF (pymupdf) — 3 sayfa, gercek metin katmani
import pymupdf  # noqa: E402

doc = pymupdf.open()
for i in range(3):
    page = doc.new_page()
    page.insert_text((72, 100), f"[TEST] Sayfa {i + 1}", fontsize=18)
    page.insert_text((72, 140), TEXT[:180], fontsize=10)
    page.insert_text((72, 170), TEXT[180:], fontsize=10)
p = OUT / "test_kitap.pdf"
doc.save(p)
doc.close()
made.append(str(p))

# ── Taranmis (metin katmani OLMAYAN) PDF — yalnizca goruntu
from PIL import Image, ImageDraw  # noqa: E402

img = Image.new("RGB", (1240, 1754), "white")
d = ImageDraw.Draw(img)
d.text((80, 120), "[TEST] TARANMIS SAYFA - metin katmani yok", fill="black")
scan_png = OUT / "_scan.png"
img.save(scan_png)
doc = pymupdf.open()
page = doc.new_page(width=595, height=842)
page.insert_image(pymupdf.Rect(0, 0, 595, 842), filename=str(scan_png))
p = OUT / "test_taranmis.pdf"
doc.save(p)
doc.close()
made.append(str(p))

# ── PPTX (python-pptx) — 3 slayt
from pptx import Presentation  # noqa: E402
from pptx.util import Pt  # noqa: E402

prs = Presentation()
for i in range(3):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = f"[TEST] Slayt {i + 1}"
    body = slide.placeholders[1].text_frame
    body.text = TEXT[:120]
    body.paragraphs[0].font.size = Pt(14)
    body.add_paragraph().text = TEXT[120:240]
p = OUT / "test_sunum.pptx"
prs.save(p)
made.append(str(p))

# ── DOCX (python-docx)
from docx import Document  # noqa: E402

document = Document()
document.add_heading("[TEST] Ders Notu", level=1)
document.add_paragraph(TEXT)
document.add_heading("[TEST] Kazanimlar", level=2)
for k in ("Piaget evrelerini siralar", "ZPD kavramini aciklar", "Baglanma tiplerini ayirt eder"):
    document.add_paragraph(k, style="List Bullet")
p = OUT / "test_belge.docx"
document.save(p)
made.append(str(p))

# ── EPUB (elle, minimal EPUB 3 paketi — ebooklib kurulu degil)
epub = OUT / "test_kitap.epub"
with zipfile.ZipFile(epub, "w") as z:
    z.writestr("mimetype", "application/epub+zip", zipfile.ZIP_STORED)
    z.writestr(
        "META-INF/container.xml",
        '<?xml version="1.0"?><container version="1.0" '
        'xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
        '<rootfiles><rootfile full-path="OEBPS/content.opf" '
        'media-type="application/oebps-package+xml"/></rootfiles></container>',
    )
    z.writestr(
        "OEBPS/content.opf",
        '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" '
        'unique-identifier="bid"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
        "<dc:identifier id=\"bid\">test-001</dc:identifier><dc:title>[TEST] EPUB Kitap</dc:title>"
        "<dc:language>tr</dc:language></metadata><manifest>"
        '<item id="c1" href="ch1.xhtml" media-type="application/xhtml+xml"/>'
        '<item id="c2" href="ch2.xhtml" media-type="application/xhtml+xml"/>'
        '</manifest><spine><itemref idref="c1"/><itemref idref="c2"/></spine></package>',
    )
    for n, body in ((1, TEXT[:200]), (2, TEXT[200:])):
        z.writestr(
            f"OEBPS/ch{n}.xhtml",
            '<?xml version="1.0" encoding="utf-8"?>'
            '<html xmlns="http://www.w3.org/1999/xhtml"><head><title>'
            f"Bolum {n}</title></head><body><h1>[TEST] Bolum {n}</h1><p>{body}</p></body></html>",
        )
made.append(str(epub))

# ── Gorseller
img = Image.new("RGB", (1000, 700), "white")
d = ImageDraw.Draw(img)
d.text((40, 60), "[TEST] TAHTA NOTU", fill="black")
d.text((40, 120), "Piaget: duyusal-motor -> soyut islemler", fill="black")
p = OUT / "test_gorsel.png"
img.save(p)
made.append(str(p))
p = OUT / "test_gorsel.jpg"
img.convert("RGB").save(p, quality=88)
made.append(str(p))

# ── Duz metin (.txt) — ALLOWED_EXTENSIONS'ta HICBIR tur icin yok, 422 beklenir
p = OUT / "test_metin.txt"
p.write_text(TEXT, encoding="utf-8")
made.append(str(p))

# ── Bos dosya (0 bayt) — 422 beklenir
p = OUT / "test_bos.pdf"
p.write_bytes(b"")
made.append(str(p))

# ── Uzantisi PDF ama icerigi PDF olmayan dosya (bozuk) — indeksleme hatasi beklenir
p = OUT / "test_bozuk.pdf"
p.write_bytes(b"bu bir PDF degil, sadece duz metin\n" * 20)
made.append(str(p))

scan_png.unlink(missing_ok=True)
for m in made:
    print(f"{Path(m).stat().st_size:>9} {m}")
