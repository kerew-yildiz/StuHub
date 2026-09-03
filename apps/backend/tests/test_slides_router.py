"""Guide slides router testleri (Faz 2.1)."""

import io
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

PPTX_MIME = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)
PPT_MIME = "application/vnd.ms-powerpoint"
LEGACY_PPT_PATH = Path(__file__).parent / "fixtures" / "legacy_presentation.ppt"


def _make_pptx_bytes() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # yalnız başlık düzeni
    title = slide.shapes.title
    if title is not None and title.has_text_frame:
        title.text_frame.text = "Bağlı Listeler"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(4))
    box.text_frame.text = "Düğümler ve işaretçiler."
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


async def _make_course_with_chapter(client) -> int:
    resp = await client.post("/api/terms", json={"name": "2026 Bahar"})
    term_id = resp.json()["id"]
    resp = await client.post(f"/api/terms/{term_id}/courses", json={"name": "Veri Yapıları"})
    course_id = resp.json()["id"]
    resp = await client.post(f"/api/courses/{course_id}/chapters", json={"title": "Bağlı Listeler"})
    return resp.json()["id"]


async def test_upload_and_list_slides(client):
    chapter_id = await _make_course_with_chapter(client)
    resp = await client.post(
        f"/api/chapters/{chapter_id}/slides",
        files={"file": ("sunum.pptx", _make_pptx_bytes(), PPTX_MIME)},
    )
    assert resp.status_code == 201
    slides = resp.json()
    assert len(slides) == 1
    assert slides[0]["slide_no"] == 1
    assert "Bağlı Listeler" in slides[0]["content_text"]
    assert slides[0]["material_id"] is not None

    resp = await client.get(f"/api/chapters/{chapter_id}/slides")
    assert resp.status_code == 200
    assert len(resp.json()) == 1


async def test_upload_pdf_slides(client):
    import pymupdf

    font_file = r"C:\Windows\Fonts\arial.ttf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text(
        (72, 72), "PDF sunum sayfası", fontsize=11, fontname="arial", fontfile=font_file
    )
    buf = io.BytesIO()
    doc.save(buf)
    doc.close()

    chapter_id = await _make_course_with_chapter(client)
    resp = await client.post(
        f"/api/chapters/{chapter_id}/slides",
        files={"file": ("sunum.pdf", buf.getvalue(), "application/pdf")},
    )
    assert resp.status_code == 201
    slides = resp.json()
    assert len(slides) == 1
    assert "PDF sunum sayfası" in slides[0]["content_text"]


async def test_slides_validation(client):
    chapter_id = await _make_course_with_chapter(client)
    # yanlış uzantı
    resp = await client.post(
        f"/api/chapters/{chapter_id}/slides",
        files={"file": ("notlar.txt", b"merhaba", "text/plain")},
    )
    assert resp.status_code == 422

    # var olmayan chapter
    resp = await client.post(
        "/api/chapters/9999/slides",
        files={"file": ("sunum.pdf", b"%PDF-1.4", "application/pdf")},
    )
    assert resp.status_code == 404


async def test_second_deck_continues_slide_numbers(client):
    chapter_id = await _make_course_with_chapter(client)
    first = await client.post(
        f"/api/chapters/{chapter_id}/slides",
        files={"file": ("a.pptx", _make_pptx_bytes(), PPTX_MIME)},
    )
    assert first.status_code == 201
    assert first.json()[0]["slide_no"] == 1

    second = await client.post(
        f"/api/chapters/{chapter_id}/slides",
        files={"file": ("b.pptx", _make_pptx_bytes(), PPTX_MIME)},
    )
    assert second.status_code == 201
    assert second.json()[0]["slide_no"] == 2

    resp = await client.get(f"/api/chapters/{chapter_id}/slides")
    assert [s["slide_no"] for s in resp.json()] == [1, 2]


async def test_delete_slide(client):
    chapter_id = await _make_course_with_chapter(client)
    created = await client.post(
        f"/api/chapters/{chapter_id}/slides",
        files={"file": ("sunum.pptx", _make_pptx_bytes(), PPTX_MIME)},
    )
    slide_id = created.json()[0]["id"]
    resp = await client.delete(f"/api/slides/{slide_id}")
    assert resp.status_code == 204
    resp = await client.get(f"/api/chapters/{chapter_id}/slides")
    assert resp.json() == []


async def test_upload_legacy_ppt_slides(client):
    """Eski ikili .ppt, ppt2pptx ile .pptx'e çevrilip aynı yoldan okunmalı (bkz.
    services/slides_service.py `_ensure_pptx`)."""
    chapter_id = await _make_course_with_chapter(client)
    resp = await client.post(
        f"/api/chapters/{chapter_id}/slides",
        files={"file": ("eski_sunum.ppt", LEGACY_PPT_PATH.read_bytes(), PPT_MIME)},
    )
    assert resp.status_code == 201
    slides = resp.json()
    assert len(slides) == 2
    assert "ppt2pptx visual fixture" in slides[0]["content_text"]
    assert "second slide" in slides[1]["content_text"]
