"""PPTX metin çıkarımı testleri (Yetenek 01)."""

from pptx import Presentation
from pptx.util import Inches

from src.services.slides_service import SlidesError, extract_pptx_slides, find_soffice


def _set_slide_title(slide, text: str) -> None:
    title = slide.shapes.title
    if title is not None and title.has_text_frame:
        title.text_frame.text = text


def _set_slide_body(slide, text: str) -> None:
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(4))
    box.text_frame.text = text


def _make_pptx(path):
    prs = Presentation()
    slide1 = prs.slides.add_slide(prs.slide_layouts[5])  # yalnız başlık düzeni
    _set_slide_title(slide1, "Giriş")
    _set_slide_body(slide1, "Bu dersin konusu.")
    notes_frame = slide1.notes_slide.notes_text_frame
    if notes_frame is not None:
        notes_frame.text = "Konuşmacı notu"

    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    _set_slide_title(slide2, "Sonuç")
    prs.save(str(path))


def test_extract_pptx_slides(tmp_path):
    pptx = tmp_path / "sunum.pptx"
    _make_pptx(pptx)
    slides = extract_pptx_slides(str(pptx))
    assert len(slides) == 2
    assert "Giriş" in slides[0]["text"]
    assert "Bu dersin konusu." in slides[0]["text"]
    assert "[Not] Konuşmacı notu" in slides[0]["text"]
    assert slides[1]["text"] == "Sonuç"


def test_corrupt_pptx_error(tmp_path):
    bad = tmp_path / "bozuk.pptx"
    bad.write_bytes(b"zip degil")
    try:
        extract_pptx_slides(str(bad))
        raise AssertionError("bozuk PPTX hatasız geçmemeli")
    except SlidesError:
        pass


def test_find_soffice_returns_none_or_path():
    # Kurulu olmayabilir; dönerse ikili adı dolu olmalı
    soffice = find_soffice()
    if soffice is not None:
        assert soffice
