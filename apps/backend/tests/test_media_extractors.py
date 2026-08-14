"""Medya çıkarım paketi testleri (Yetenek 11) — gerçek model indirme YOK (mock'lu).

Mock stratejisi: yt-dlp/whisper/rapidocr ağ ve model işleri sahte nesnelerle taklit
edilir; DOCX/EPUB gerçek dosya inşasıyla, metin/segment/timestamp mantığı gerçekten
doğrulanır.
"""

from __future__ import annotations

import urllib.request
import zipfile
from pathlib import Path

import faster_whisper
import pytest
import rapidocr_onnxruntime
import yt_dlp
from docx import Document

from src.config import settings
from src.services.media_extractors import (
    EXTRACTORS,
    extract_for,
    transcript_to_text,
)
from src.services.media_extractors import audio as audio_extractor
from src.services.media_extractors import docx as docx_extractor
from src.services.media_extractors import epub as epub_extractor
from src.services.media_extractors import ocr as ocr_extractor
from src.services.media_extractors import text as text_extractor
from src.services.media_extractors import youtube as youtube_extractor


class _FakeResponse:
    def __init__(self, data: bytes):
        self._data = data

    def read(self) -> bytes:
        return self._data

    def __enter__(self):
        return self

    def __exit__(self, *args):
        return False


class _FakeSegment:
    def __init__(self, start: float, end: float, text: str):
        self.start = start
        self.end = end
        self.text = text


class _FakeWhisperModel:
    def __init__(self, *args, **kwargs):
        pass

    def transcribe(self, path):
        return (
            [_FakeSegment(0.0, 1.5, "merhaba"), _FakeSegment(1.5, 3.0, "dünya")],
            {},
        )


def _make_docx(path, paragraphs=(), table_rows=()):
    doc = Document()
    for para in paragraphs:
        doc.add_paragraph(para)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_idx, row_values in enumerate(table_rows):
            for col_idx, value in enumerate(row_values):
                table.rows[row_idx].cells[col_idx].text = value
    doc.save(str(path))


def _make_epub(path, chapters):
    container = (
        '<?xml version="1.0"?>'
        '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
        '<rootfiles><rootfile full-path="OEBPS/content.opf" '
        'media-type="application/oebps-package+xml"/></rootfiles></container>'
    )
    items = "".join(
        f'<item id="c{i}" href="chapter{i}.xhtml" media-type="application/xhtml+xml"/>'
        for i in range(1, len(chapters) + 1)
    )
    spine = "".join(f'<itemref idref="c{i}"/>' for i in range(1, len(chapters) + 1))
    opf = (
        '<?xml version="1.0"?>'
        '<package xmlns="http://www.idpf.org/2007/opf" version="3.0">'
        f"<manifest>{items}</manifest><spine>{spine}</spine></package>"
    )
    with zipfile.ZipFile(str(path), "w") as zf:
        zf.writestr("META-INF/container.xml", container)
        zf.writestr("OEBPS/content.opf", opf)
        for i, chapter in enumerate(chapters, start=1):
            zf.writestr(f"OEBPS/chapter{i}.xhtml", chapter)


# ── text ───────────────────────────────────────────────────────────────────────


def test_text_empty_raises():
    with pytest.raises(RuntimeError):
        text_extractor.extract("")
    with pytest.raises(RuntimeError):
        text_extractor.extract("   ")


def test_text_single_short_segment():
    segments = text_extractor.extract("Kısa metin.")
    assert segments == [{"segment": 0, "text": "Kısa metin.", "start": None, "end": None}]


def test_text_splits_long_text_at_word_boundary():
    word = "kelime"
    text = " ".join([word] * 400)  # ~2800 karakter > 1500
    segments = text_extractor.extract(text)
    assert len(segments) >= 2
    assert [s["segment"] for s in segments] == list(range(len(segments)))
    assert all(len(s["text"]) <= 1500 + 100 for s in segments)
    for seg in segments:
        assert seg["text"] == seg["text"].strip()


# ── docx ───────────────────────────────────────────────────────────────────────


def test_docx_extracts_paragraphs_and_tables(tmp_path):
    path = tmp_path / "belge.docx"
    _make_docx(
        path,
        paragraphs=["Merhaba dünya", "İkinci paragraf"],
        table_rows=[["A1", "B1"], ["A2", "B2"]],
    )
    segments = docx_extractor.extract(str(path))
    joined = " ".join(s["text"] for s in segments)
    assert "Merhaba dünya" in joined
    assert "İkinci paragraf" in joined
    assert "A1" in joined and "B2" in joined
    assert all(s["start"] is None and s["end"] is None for s in segments)


def test_docx_empty_raises(tmp_path):
    path = tmp_path / "bos.docx"
    Document().save(str(path))
    with pytest.raises(RuntimeError):
        docx_extractor.extract(str(path))


def test_docx_missing_file_raises(tmp_path):
    with pytest.raises(RuntimeError):
        docx_extractor.extract(str(tmp_path / "yok.docx"))


# ── epub ───────────────────────────────────────────────────────────────────────


def test_epub_extracts_xhtml_in_order(tmp_path):
    path = tmp_path / "kitap.epub"
    _make_epub(
        path,
        [
            "<html><head><style>.x{color:red}</style></head><body>"
            "<h1>Başlık</h1><p>Merhaba &amp; dünya</p></body></html>",
            "<html><body><p>İkinci bölüm içeriği</p></body></html>",
        ],
    )
    segments = epub_extractor.extract(str(path))
    joined = "\n".join(s["text"] for s in segments)
    assert "Merhaba & dünya" in joined
    assert "İkinci bölüm içeriği" in joined
    assert "color" not in joined  # <style> içeriği metne karışmamalı
    assert all(s["start"] is None and s["end"] is None for s in segments)


def test_epub_empty_raises(tmp_path):
    path = tmp_path / "bos.epub"
    _make_epub(path, ["<html><body><p>   </p></body></html>"])
    with pytest.raises(RuntimeError):
        epub_extractor.extract(str(path))


def test_epub_missing_file_raises(tmp_path):
    with pytest.raises(ValueError):
        epub_extractor.extract(str(tmp_path / "yok.epub"))


# ── youtube ────────────────────────────────────────────────────────────────────


def test_youtube_captions_parsed_to_segments(monkeypatch):
    info = {"subtitles": {"tr": [{"url": "http://fake/captions.vtt"}]}, "automatic_captions": {}}
    monkeypatch.setattr(
        yt_dlp.YoutubeDL, "extract_info", lambda self, url, download=False: info
    )
    vtt = (
        "WEBVTT\n\n"
        "00:00:00.000 --> 00:00:02.500\n"
        "Merhaba dünya\n\n"
        "00:00:02.500 --> 00:00:05.000\n"
        "İkinci satır\n"
    )
    monkeypatch.setattr(
        urllib.request, "urlopen", lambda url, timeout=30: _FakeResponse(vtt.encode())
    )
    segments = youtube_extractor.extract("https://youtube.com/watch?v=abc")
    assert len(segments) == 2
    assert segments[0]["text"] == "Merhaba dünya"
    assert segments[0]["start"] == 0.0
    assert segments[0]["end"] == 2.5
    assert segments[1]["text"] == "İkinci satır"
    assert segments[1]["start"] == 2.5
    assert segments[1]["end"] == 5.0


def test_youtube_fallback_whisper(monkeypatch):
    info = {"subtitles": {}, "automatic_captions": {}}
    monkeypatch.setattr(
        yt_dlp.YoutubeDL, "extract_info", lambda self, url, download=False: info
    )

    def fake_download(self, urls):
        tmpl = self.params.get("outtmpl") or "audio"
        if isinstance(tmpl, dict):
            tmpl = tmpl.get("default") or next(iter(tmpl.values()), "audio")
        Path(str(tmpl).replace("%(ext)s", "webm")).write_bytes(b"x")
        return 0

    monkeypatch.setattr(yt_dlp.YoutubeDL, "download", fake_download)
    monkeypatch.setattr(faster_whisper, "WhisperModel", _FakeWhisperModel)

    segments = youtube_extractor.extract("https://youtube.com/watch?v=abc")
    assert [s["text"] for s in segments] == ["merhaba", "dünya"]
    assert segments[0]["start"] == 0.0
    assert segments[1]["end"] == 3.0


def test_youtube_invalid_url_raises():
    with pytest.raises(ValueError):
        youtube_extractor.extract("bu bir url degil")


# ── audio ──────────────────────────────────────────────────────────────────────


def test_audio_transcribes_with_whisper(monkeypatch, tmp_path):
    audio_file = tmp_path / "kayit.mp3"
    audio_file.write_bytes(b"dummy")
    monkeypatch.setattr(faster_whisper, "WhisperModel", _FakeWhisperModel)

    segments = audio_extractor.extract(str(audio_file))
    assert [s["text"] for s in segments] == ["merhaba", "dünya"]
    assert segments[0]["start"] == 0.0
    assert segments[1]["end"] == 3.0


def test_audio_missing_file_raises(tmp_path):
    with pytest.raises(ValueError):
        audio_extractor.extract(str(tmp_path / "yok.mp3"))


# ── ocr ────────────────────────────────────────────────────────────────────────


def test_ocr_disabled_raises(monkeypatch):
    monkeypatch.setattr(settings, "ocr_enabled", False)
    with pytest.raises(RuntimeError):
        ocr_extractor.extract("gorsel.png")


def test_ocr_extracts_text(monkeypatch, tmp_path):
    image = tmp_path / "gorsel.png"
    image.write_bytes(b"dummy")

    class FakeEngine:
        def __init__(self, *args, **kwargs):
            pass

        def __call__(self, path):
            return [[[[0, 0], [10, 0], [10, 10], [0, 10]], "Merhaba", 0.99]], 0.0

    monkeypatch.setattr(rapidocr_onnxruntime, "RapidOCR", FakeEngine)

    segments = ocr_extractor.extract(str(image))
    assert segments[0]["text"] == "Merhaba"
    assert segments[0]["start"] is None


def test_ocr_bad_extension_raises(tmp_path):
    bad = tmp_path / "dosya.txt"
    bad.write_bytes(b"x")
    with pytest.raises(ValueError):
        ocr_extractor.extract(str(bad))


# ── transcript_to_text + dispatcher ────────────────────────────────────────────


def test_transcript_to_text_formats_timestamps():
    segments = [
        {"segment": 0, "text": "İlk satır", "start": 0.0, "end": 12.5},
        {"segment": 1, "text": "İkinci satır", "start": 65.0, "end": 80.0},
        {"segment": 2, "text": "Damgasız", "start": None, "end": None},
    ]
    assert transcript_to_text(segments) == "[00:00] İlk satır\n[01:05] İkinci satır\nDamgasız"


def test_transcript_to_text_skips_empty():
    segments = [{"segment": 0, "text": "   ", "start": 0.0, "end": 1.0}]
    assert transcript_to_text(segments) == ""


def test_extract_for_dispatches_and_rejects_unknown():
    assert set(EXTRACTORS) == {"youtube", "audio", "docx", "epub", "image", "text"}
    with pytest.raises(ValueError):
        extract_for("bilinmeyen", "kaynak")


def test_extract_for_text_dispatches():
    segments = extract_for("text", "Merhaba")
    assert segments[0]["text"] == "Merhaba"
