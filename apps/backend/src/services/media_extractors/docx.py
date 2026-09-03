"""DOCX metin çıkarımı — python-docx (tembel import, Yetenek 11)."""

from __future__ import annotations

from src.services.media_extractors import split_text_segments


def extract(path: str) -> list[dict]:
    """Paragrafları + tablo hücrelerini toplar; ~1500 karakterlik segmentlere böler."""
    from docx import Document

    try:
        document = Document(path)
    except Exception as exc:
        raise RuntimeError("DOCX okunamadı, dosya bozuk olabilir.") from exc

    lines: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            if line:
                lines.append(line)

    full_text = "\n".join(lines).strip()
    if not full_text:
        raise RuntimeError("DOCX'ten metin çıkarılamadı.")
    parts = split_text_segments(full_text)
    return [
        {"segment": i, "text": part, "start": None, "end": None}
        for i, part in enumerate(parts)
    ]
