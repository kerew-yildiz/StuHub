"""Alan 9 test dosyalari uretici (gecici, temizlik listesinde)."""
import pathlib
import pymupdf
from pptx import Presentation
from pptx.util import Inches, Pt

OUT = pathlib.Path(__file__).parent

# 1) Duz ASCII 3 sayfalik PDF
doc = pymupdf.open()
for i in range(1, 4):
    p = doc.new_page()
    p.insert_text((72, 100), f"TEST Chapter Page {i}", fontsize=20)
    p.insert_text((72, 140), f"This is sample body text for page {i} of the Alan09 upload test.", fontsize=11)
    p.insert_text((72, 165), "Cognitive development proceeds through discrete stages.", fontsize=11)
doc.save(OUT / "alan09_small.pdf")
doc.close()

# 2) Turkce karakterli 2 sayfalik PDF (helv Latin-1 disi glifleri kaldirmaz; TR icin
#    gomulu bir Unicode font sart -> pymupdf'in dahili 'china-ss' benzeri degil,
#    Windows'ta bulunan Arial kullanilir)
TR1 = "Gelişimsel Psikoloji — Çocukluk Çağı Özellikleri"
TR2 = "İğne, ışık, şğüöçİĞÜŞÖÇ; bilişsel gelişim aşamaları çok önemlidir."
doc = pymupdf.open()
for i, head in enumerate([TR1, "Öğrenme ve Bellek Üzerine Değerlendirme"], start=1):
    p = doc.new_page()
    p.insert_text((60, 100), head, fontsize=18, fontfile=r"C:\Windows\Fonts\arial.ttf", fontname="AR")
    p.insert_text((60, 140), TR2, fontsize=11, fontfile=r"C:\Windows\Fonts\arial.ttf", fontname="AR")
doc.save(OUT / "alan09_turkce.pdf")
doc.close()

# 3) 3 slaytlik PPTX
prs = Presentation()
blank = prs.slide_layouts[6]
for i, (title, body) in enumerate([
    ("Slayt Bir: Giris", "Ilk slaydin govde metni burada."),
    ("Slayt Iki: Gelisim", "Ikinci slaydin govde metni, bilissel gelisim."),
    ("Slayt Uc: Sonuc", "Ucuncu slaydin govde metni ve ozet."),
], start=1):
    s = prs.slides.add_slide(blank)
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    tb.text_frame.text = title
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    bb = s.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(2))
    bb.text_frame.text = body
prs.save(OUT / "alan09_slides.pptx")

for f in sorted(OUT.glob("alan09_*")):
    print(f.name, f.stat().st_size)
print("pages_small", pymupdf.open(OUT / "alan09_small.pdf").page_count)
print("pages_tr", pymupdf.open(OUT / "alan09_turkce.pdf").page_count)
